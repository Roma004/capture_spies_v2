# %%
from openpyxl import load_workbook
from datetime import date
import pandas as pd

inp_dir = "../../data/tmp/YourBoardingPassDotAero/"
res_dir = "../../data/unified/boarding-pass/"

# %%
def get_inp_filename(dt: date) -> str:
    return f"{inp_dir}YourBoardingPassDotAero-{dt}.xlsx"

def xlsx_to_dict(file_path, data_only=True):
    """
    Читает все листы xlsx-файла и возвращает словарь:
      { 'имя_листа': [[строка1], [строка2], ...] }
    
    :param file_path: путь к xlsx-файлу
    :param data_only: если True — возвращать значения формул, а не сами формулы
    """
    wb = load_workbook(file_path, data_only=data_only)
    result = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        # values_only=True сразу отдаёт значения ячеек
        result[sheet_name] = [list(row) for row in ws.iter_rows(values_only=True)]
    wb.close()
    return result


# %%

cur_date = date(2017, 1, 1)
dfs = []
for _ in range(365):
    data = xlsx_to_dict(get_inp_filename(cur_date))
    d = data
    sheet_data_list = []
    for sheet, rows in d.items():
        # удалить каждую вторую строку в d
        d[sheet] = [row for i, row in enumerate(rows) if i % 2 == 0]

    for sheet, rows in d.items():
        """
        извлечь всю полезную инфлормацию из листа
        Example sheet data:

        ['BOARDING PASS', None, None, None, None, 'SEQUENCE:', None, 32]
        ['MRS', 'LIDIYA ZHDANOVA', None, None, None, None, None, 'Y']
        ['SU5436', None, None, 'VLADIVOSTOK', None, None, None, 'SEOUL']
        ['GATE', 'N/A', None, 'VVO', None, '->', None, 'REA']
        ['2017-01-01', '/', '11:25', None, 'Operated by Some Other Airline', None, None, None]
        ['Boarding is ended 20 minutes before departure time', None, None, None, None, None, 'SEAT', 'N/A']
        ['PNR:', 'DYELAO', None, 'E-TICKET', '7100246988860481', None, None, None]
        """

        sheet_data = {
            "sheet": sheet,
            "BOARDING_PASS": rows[0][0],
            "PASSENGER_NAME": rows[1][1],
            "FLIGHT_NUMBER": rows[2][0],
            "FROM": rows[2][3],
            "TO": rows[2][7],
            "GATE": rows[3][1],
            "DEPARTURE_AIRPORT": rows[3][3],
            "ARRIVAL_AIRPORT": rows[3][7],
            "DEPARTURE_DATE": rows[4][0],
            "DEPARTURE_TIME": rows[4][2],
            "OPERATED_BY": rows[4][4],
            "SEAT": rows[5][7],
            "PNR": rows[6][1],
            "E_TICKET": rows[6][4],
        }

        assert sheet_data["BOARDING_PASS"] == "BOARDING PASS", f"Unexpected value in BOARDING_PASS: {sheet_data['BOARDING_PASS']}"
        sheet_data_list.append(sheet_data)
    dfs.append(pd.DataFrame(sheet_data_list))
    next_date = cur_date + pd.Timedelta(days=1)
    print(f"processed {cur_date}")
    if next_date.month != cur_date.month:
        df = pd.concat(dfs, ignore_index=True)
        df.to_csv(f"{res_dir}{cur_date.year}-{cur_date.month:02d}.csv", index=False)
        print(f"processed {cur_date.year}-{cur_date.month:02d}.csv")
        del df
        del dfs
        dfs = []
    cur_date = next_date
